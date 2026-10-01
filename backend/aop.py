"""AOP (Agent Operating Procedure, .docx) awareness for chat.

Off by default -- only used when a video has an AOP uploaded AND the user has
switched "AOP mode" on for that chat. Normal chat (chat.py) is untouched and
stays cheap; this module only runs when explicitly asked for.

The AOP content itself always comes from whatever .docx the user uploads --
nothing here generates or fabricates AOP structure or content.

Same right-sized-for-scale note as chat.py: AOP section content is stuffed
into the prompt directly (truncated per-section as a cost guard), not
retrieved via embeddings. Fine for a single procedure document; would need
real chunking if AOPs grow very large.
"""

import json
import time
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from google import genai
from google.genai import types

from . import usage as usage_module
from .ingestion.reslog import LOG_PATH
from .settings_store import get_gemini_api_key, get_gemini_model
from .text_corrections import repair_flattened_bullets

MAX_PARAGRAPH_LEVELS = 4
SECTION_CONTENT_CHAR_CAP = 1500  # per-section cost guard when building the prompt


def _iter_block_items(doc: Document):
    """Paragraphs AND tables, in true document order -- python-docx's own
    `doc.paragraphs`/`doc.tables` are separate lists with no ordering between
    them, which silently drops tables from a straight paragraph walk. Real
    AOPs put a lot of structured content in tables, so this matters."""
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "para", Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, doc)


def _iter_sections(doc: Document) -> list[dict]:
    """Flat list of sections. Each keeps live references (the heading
    paragraph, and every content paragraph) into the *same* document instance
    -- needed so apply_suggestion() can add/remove/replace real content after
    re-opening and re-walking the doc the same way."""
    sections = []
    counters = [0] * MAX_PARAGRAPH_LEVELS
    current = None
    for kind, item in _iter_block_items(doc):
        if kind == "para":
            style = item.style.name or ""
            if style.startswith("Heading"):
                try:
                    level = min(int(style.replace("Heading", "").strip() or 1), MAX_PARAGRAPH_LEVELS)
                except ValueError:
                    level = 1
                counters[level - 1] += 1
                for i in range(level, MAX_PARAGRAPH_LEVELS):
                    counters[i] = 0
                section_id = ".".join(str(c) for c in counters[:level] if c) or str(len(sections) + 1)
                current = {
                    "id": section_id, "title": item.text.strip(), "level": level,
                    "heading_para": item, "content_paras": [], "table_rows": [],
                }
                sections.append(current)
            else:
                if item.text.strip() and current:
                    current["content_paras"].append(item)
        elif kind == "table" and current:
            for row in item.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells)
                if row_text.replace("|", "").strip():
                    current["table_rows"].append(row_text)
    return sections


def parse_structure(docx_path: Path) -> list[dict]:
    doc = Document(str(docx_path))
    sections = _iter_sections(doc)
    return [
        {
            "id": s["id"], "title": s["title"], "level": s["level"],
            "paragraphs": [p.text.strip() for p in s["content_paras"] if p.text.strip()],
            "tables": s["table_rows"],
        }
        for s in sections
    ]


def apply_suggestion(docx_path: Path, section_id: str, action: str, text: str = "", existing_text: str = "") -> None:
    """action: "add" (append text at end of section), "remove" (delete the
    content paragraph whose text exactly matches existing_text), or "replace"
    (remove that paragraph, insert text in roughly its place).

    Every inserted paragraph is italicized so a human reviewer can see at a
    glance what was AI-added/changed. Always edits the *working* copy --
    aop_original.docx (written once, at upload time) is the untouched backup.
    Table rows are not editable via remove/replace in this version.
    """
    doc = Document(str(docx_path))
    sections = _iter_sections(doc)
    idx = next((i for i, s in enumerate(sections) if s["id"] == section_id), None)
    if idx is None:
        raise ValueError(f"no such AOP section: {section_id!r}")
    section = sections[idx]
    next_section = sections[idx + 1] if idx + 1 < len(sections) else None

    removed_next_sibling = None
    if action in ("remove", "replace"):
        if not existing_text.strip():
            raise ValueError("remove/replace requires existing_text to match against")
        target = next((p for p in section["content_paras"] if p.text.strip() == existing_text.strip()), None)
        if target is None:
            raise ValueError("couldn't find that exact existing text in this section -- it may have already changed")
        removed_next_sibling = target._element.getnext()
        target._element.getparent().remove(target._element)

    if action == "add" or (action == "replace" and text.strip()):
        if action == "replace" and isinstance(removed_next_sibling, CT_P):
            new_para = Paragraph(removed_next_sibling, doc).insert_paragraph_before()
        elif next_section:
            new_para = next_section["heading_para"].insert_paragraph_before()
        else:
            new_para = doc.add_paragraph()
        run = new_para.add_run(text)
        run.italic = True
    elif action not in ("add", "remove", "replace"):
        raise ValueError(f"unknown action: {action!r}")

    doc.save(str(docx_path))


AOP_SCHEMA = {
    "type": "object",
    "properties": {
        "answer": {"type": "string"},
        "citations": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "node_id": {"type": "string"},
                    "label": {"type": "string"},
                    "t_ms": {"type": "integer"},
                },
                "required": ["node_id", "t_ms"],
            },
        },
        "aop_suggestions": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "section_id": {"type": "string"},
                    "section_title": {"type": "string"},
                    "action": {"type": "string", "enum": ["add", "remove", "replace"]},
                    "existing_text": {
                        "type": "string",
                        "description": "For remove/replace only: the exact existing sentence copied verbatim from that section's content above.",
                    },
                    "suggestion": {
                        "type": "string",
                        "description": "For add/replace only: the new text to insert.",
                    },
                },
                "required": ["section_id", "action"],
            },
        },
    },
    "required": ["answer", "citations", "aop_suggestions"],
}

AOP_PROMPT = """You are helping a forward-deployed engineer keep an Agent Operating Procedure (AOP) document in sync while they review a process-mapping video. The AOP is what actually gets used to build an automation agent, so gaps and errors here matter.

You have three sources of truth:
1. The video's knowledge graph (steps, systems, actors, data, decisions -- each with timestamps)
2. The video's transcript
3. The current AOP's sections (id, title, existing paragraph content, and any table rows) -- this is a REAL document the user uploaded; treat its existing content as ground truth for what the AOP currently says.

First, answer the user's question using the graph + transcript, with citations, exactly as you normally would. Only cite the nodes you actually reference in the answer -- do not dump every matching node into citations. Format the answer as plain Markdown, kept minimal: short paragraphs, and a "- " bulleted list instead of one long comma-separated sentence whenever you're listing three or more items. `answer` is a JSON string, but it must still contain real newline characters: a blank line between paragraphs, and every "- " bullet starting on its own new line -- never run bullets together separated only by a space, or they will render as one wall of text instead of a list.

Then compare what the video shows against what the AOP currently documents, and propose concrete edits via aop_suggestions:
- action "add": the video shows something relevant to this question that the AOP is missing entirely. Give the exact section id and the new text.
- action "remove": something existing in the AOP contradicts or is clearly superseded by what the video shows. Copy the exact existing sentence into existing_text, verbatim, so it can be located and deleted.
- action "replace": an existing AOP statement is wrong or outdated per the video. Copy the exact existing sentence into existing_text AND give the corrected text in suggestion.

Only propose a change the video/graph actually supports -- an empty aop_suggestions list is correct when nothing is missing or wrong for this question. Never propose vague notes like "consider adding detail"; word suggestions so they could be dropped straight into the document.

AOP sections:
{aop_json}

Knowledge graph:
{graph_json}

Transcript:
{transcript_text}

Conversation so far:
{history}

User question: {question}
"""


def _log_usage(label: str, model: str, usage, elapsed_s: float) -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with LOG_PATH.open("a") as f:
        f.write(
            f"| {time.strftime('%Y-%m-%d')} | {label} | Gemini AOP-chat call (`{model}`) "
            f"| prompt={usage.prompt_token_count} candidates={usage.candidates_token_count} "
            f"thoughts={usage.thoughts_token_count or 0} total={usage.total_token_count} tokens "
            f"| {elapsed_s:.1f}s wall | network call -- logged for free-tier budget tracking |\n"
        )


def answer_with_aop(work_dir: Path, question: str, history: list[dict]) -> dict:
    graph = json.loads((work_dir / "graph.json").read_text())
    transcript = json.loads((work_dir / "transcript.json").read_text())
    aop_structure = json.loads((work_dir / "aop_structure.json").read_text())
    transcript_text = "\n".join(f"[{s['start_ms']}-{s['end_ms']}ms] {s['text']}" for s in transcript)
    history_text = "\n".join(f"{m['role']}: {m['content']}" for m in history[-6:]) or "(none yet)"

    api_key = get_gemini_api_key()
    if not api_key:
        raise RuntimeError("No Gemini API key configured. Add one in Settings.")
    model = get_gemini_model()
    client = genai.Client(api_key=api_key)

    aop_compact = []
    for s in aop_structure:
        content = " ".join(s.get("paragraphs", []))
        if s.get("tables"):
            content += " TABLES: " + " ; ".join(s["tables"])
        aop_compact.append({"id": s["id"], "title": s["title"], "content": content[:SECTION_CONTENT_CHAR_CAP]})

    aop_json = json.dumps(aop_compact)[:60_000]
    graph_json = json.dumps(graph)[:60_000]
    transcript_capped = transcript_text[:60_000]
    prompt = AOP_PROMPT.format(
        aop_json=aop_json,
        graph_json=graph_json,
        transcript_text=transcript_capped,
        history=history_text,
        question=question,
    )

    video_key = f"{work_dir.parent.name}/{work_dir.name}"
    components = usage_module.count_tokens(client, model, {
        "AOP sections": aop_json,
        "Knowledge graph": graph_json,
        "Transcript": transcript_capped,
        "Conversation history": history_text,
        "Your question": question,
    })

    t0 = time.time()
    response = client.models.generate_content(
        model=model,
        contents=[prompt],
        config=types.GenerateContentConfig(response_mime_type="application/json", response_schema=AOP_SCHEMA),
    )
    elapsed = time.time() - t0
    _log_usage(f"aop-chat:{video_key}", model, response.usage_metadata, elapsed)
    usage_module.record_call("video", video_key, model, response.usage_metadata, elapsed, components, label="AOP Chat")
    data = json.loads(response.text)
    data["answer"] = repair_flattened_bullets(data.get("answer", ""))
    return data
