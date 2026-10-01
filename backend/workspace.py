"""Workspace: upload reference docs (AOP/SOP/checklists, or platform/architecture
docs), chat grounded only in those docs, and draft a Notes document from the
conversation.

Two independent "boards" share this exact module -- "process" (understand how
the current system/process works, from AOP/SOP/intake/checklist docs) and
"agent" (understand how the agent should be built, from Zamp platform
architecture docs, diagrams, case studies). Same code, different doc set and
notes draft per board -- selected purely by the `board` argument threaded
through every function.

Supported doc formats for V1: .docx, .csv, .txt, .md. PDF is explicitly not
supported yet -- flagged, not silently dropped (see upload_doc).
"""

import csv
import io
import json
import time
import uuid
from pathlib import Path

from docx import Document
from google import genai
from google.genai import types

from . import usage as usage_module
from .extraction.schema import EDGE_TYPES, NODE_TYPES
from .ingestion import config
from .ingestion.reslog import LOG_PATH
from .settings_store import get_gemini_api_key, get_gemini_model
from .text_corrections import repair_flattened_bullets

BOARDS = {
    "process": "Process Understanding",
    "agent": "Agent Design",
}
SUPPORTED_EXTENSIONS = {".docx", ".csv", ".txt", ".md"}
DOC_TEXT_CHAR_CAP = 20_000  # per-document cost guard when building the chat prompt


def _board_dir(board: str) -> Path:
    if board not in BOARDS:
        raise ValueError(f"unknown board: {board!r}")
    d = config.DATA_DIR / "workspace" / board
    (d / "docs").mkdir(parents=True, exist_ok=True)
    (d / "threads").mkdir(parents=True, exist_ok=True)
    return d


def _index_path(board: str) -> Path:
    return _board_dir(board) / "docs_index.json"


def _load_index(board: str) -> list[dict]:
    p = _index_path(board)
    return json.loads(p.read_text()) if p.exists() else []


def _save_index(board: str, index: list[dict]) -> None:
    _index_path(board).write_text(json.dumps(index, indent=2))


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8", errors="replace") as f:
            rows = list(csv.reader(f))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"unsupported file type: {suffix!r} (supported: {sorted(SUPPORTED_EXTENSIONS)})")


def upload_doc(board: str, filename: str, content: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"'{suffix}' isn't supported yet (supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}). "
            f"PDF support hasn't been built -- ask for it if you need it."
        )
    doc_id = uuid.uuid4().hex[:10]
    dest = _board_dir(board) / "docs" / f"{doc_id}{suffix}"
    dest.write_bytes(content)

    try:
        text = _extract_text(dest)
    except Exception as e:
        dest.unlink(missing_ok=True)
        raise ValueError(f"couldn't read {filename}: {e}")

    entry = {
        "id": doc_id, "filename": filename, "chars": len(text),
        "uploaded_at": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
    index = _load_index(board)
    index.append(entry)
    _save_index(board, index)
    (_board_dir(board) / "docs" / f"{doc_id}.txt").write_text(text)
    return entry


def list_docs(board: str) -> list[dict]:
    return _load_index(board)


def delete_doc(board: str, doc_id: str) -> None:
    index = [d for d in _load_index(board) if d["id"] != doc_id]
    _save_index(board, index)
    docs_dir = _board_dir(board) / "docs"
    for p in docs_dir.glob(f"{doc_id}.*"):
        p.unlink(missing_ok=True)


def _all_docs_text(board: str) -> list[dict]:
    docs_dir = _board_dir(board) / "docs"
    out = []
    for entry in _load_index(board):
        txt_path = docs_dir / f"{entry['id']}.txt"
        if txt_path.exists():
            out.append({"filename": entry["filename"], "text": txt_path.read_text()[:DOC_TEXT_CHAR_CAP]})
    return out


# ---------------- chat threads ----------------

def _thread_path(board: str, thread_id: str) -> Path:
    return _board_dir(board) / "threads" / f"{thread_id}.json"


def list_threads(board: str) -> list[dict]:
    out = []
    for p in sorted((_board_dir(board) / "threads").glob("*.json")):
        data = json.loads(p.read_text())
        out.append({"id": data["id"], "title": data["title"], "updated_at": data["updated_at"]})
    return sorted(out, key=lambda t: t["updated_at"], reverse=True)


def create_thread(board: str, title: str = "New chat") -> dict:
    thread_id = uuid.uuid4().hex[:10]
    data = {"id": thread_id, "title": title, "messages": [], "updated_at": time.strftime("%Y-%m-%d %H:%M:%S")}
    _thread_path(board, thread_id).write_text(json.dumps(data, indent=2))
    return data


def get_thread(board: str, thread_id: str) -> dict:
    p = _thread_path(board, thread_id)
    if not p.exists():
        raise ValueError(f"no such thread: {thread_id}")
    return json.loads(p.read_text())


def _save_thread(board: str, data: dict) -> None:
    data["updated_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
    _thread_path(board, data["id"]).write_text(json.dumps(data, indent=2))


def rename_thread(board: str, thread_id: str, new_title: str) -> dict:
    thread = get_thread(board, thread_id)
    thread["title"] = new_title.strip()[:120] or thread["title"]
    _save_thread(board, thread)
    return thread


def delete_thread(board: str, thread_id: str) -> None:
    _thread_path(board, thread_id).unlink(missing_ok=True)


WORKSPACE_SCHEMA = {
    "type": "object",
    "properties": {
        "answer": {"type": "string"},
        "sources": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["answer", "sources"],
}

WORKSPACE_PROMPT = """You are answering questions using ONLY the reference documents below -- {board_label}. If the documents don't cover something, say so plainly instead of guessing or using outside knowledge.

Format the answer as plain Markdown, kept minimal: short paragraphs, and a "- " bulleted list instead of one long comma-separated sentence whenever you're listing three or more items. `answer` is a JSON string, but it must still contain real newline characters: a blank line between paragraphs, and every "- " bullet starting on its own new line -- never run bullets together separated only by a space, or they will render as one wall of text instead of a list.

List the filenames of the documents your answer actually drew from in `sources`.

Reference documents:
{docs_text}

Conversation so far:
{history}

User question: {question}
"""


def _log_usage(label: str, model: str, usage, elapsed_s: float) -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with LOG_PATH.open("a") as f:
        f.write(
            f"| {time.strftime('%Y-%m-%d')} | {label} | Gemini workspace-chat call (`{model}`) "
            f"| prompt={usage.prompt_token_count} candidates={usage.candidates_token_count} "
            f"thoughts={usage.thoughts_token_count or 0} total={usage.total_token_count} tokens "
            f"| {elapsed_s:.1f}s wall | network call -- logged for free-tier budget tracking |\n"
        )


def ask(board: str, thread_id: str, question: str) -> dict:
    thread = get_thread(board, thread_id)
    docs = _all_docs_text(board)
    if not docs:
        raise RuntimeError(f"No documents uploaded to {BOARDS[board]} yet -- upload at least one first.")

    api_key = get_gemini_api_key()
    if not api_key:
        raise RuntimeError("No Gemini API key configured. Add one in Settings.")
    model = get_gemini_model()
    client = genai.Client(api_key=api_key)

    docs_text = "\n\n".join(f"=== {d['filename']} ===\n{d['text']}" for d in docs)[:80_000]
    history_text = "\n".join(f"{m['role']}: {m['content']}" for m in thread["messages"][-8:]) or "(none yet)"

    prompt = WORKSPACE_PROMPT.format(
        board_label=BOARDS[board], docs_text=docs_text, history=history_text, question=question,
    )

    components = usage_module.count_tokens(client, model, {
        "Reference documents": docs_text,
        "Conversation history": history_text,
        "Your question": question,
    })

    t0 = time.time()
    response = client.models.generate_content(
        model=model,
        contents=[prompt],
        config=types.GenerateContentConfig(response_mime_type="application/json", response_schema=WORKSPACE_SCHEMA),
    )
    elapsed = time.time() - t0
    _log_usage(f"workspace:{board}/{thread_id}", model, response.usage_metadata, elapsed)
    usage_module.record_call("workspace", board, model, response.usage_metadata, elapsed, components, label="Chat")
    result = json.loads(response.text)
    result["answer"] = repair_flattened_bullets(result.get("answer", ""))

    thread["messages"].append({"role": "user", "content": question})
    thread["messages"].append({"role": "assistant", "content": result["answer"]})
    if thread["title"] == "New chat" and len(thread["messages"]) == 2:
        thread["title"] = question[:60]
    _save_thread(board, thread)

    return result


# ---------------- knowledge graph ----------------
# Reuses the same node/edge vocabulary as the video graph (backend/extraction/schema.py)
# so the two graph views mean the same thing by color/type. Shape differs where it has
# to: no t_start_ms/t_end_ms (no video here) -- source_docs takes its place for
# traceability, mirroring how ask() already returns `sources`.

def _graph_path(board: str) -> Path:
    return _board_dir(board) / "graph.json"


WORKSPACE_GRAPH_SCHEMA = {
    "type": "object",
    "properties": {
        "nodes": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "string", "description": "short stable kebab-case id, e.g. step-verify-lead"},
                    "type": {"type": "string", "enum": NODE_TYPES},
                    "label": {"type": "string"},
                    "description": {"type": "string"},
                    "source_docs": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["id", "type", "label", "source_docs"],
            },
        },
        "edges": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "source": {"type": "string"},
                    "target": {"type": "string"},
                    "type": {"type": "string", "enum": EDGE_TYPES},
                },
                "required": ["source", "target", "type"],
            },
        },
    },
    "required": ["nodes", "edges"],
}

WORKSPACE_GRAPH_PROMPT = """You are extracting a knowledge graph from a set of reference documents -- {board_label}.

Identify the steps, decisions, systems, actors, data entities, and artifacts these documents describe, and how they connect. Give every node a short, stable, kebab-case id, and list which document(s) it came from in source_docs (use the exact filenames given below). Connect steps in sequence with NEXT edges; use TRIGGERS/USES/DEPENDS_ON/PRODUCES/CONSUMES/ALTERNATIVE_PATH/PERFORMED_BY only where clearly supported by the text. Prefer fewer, well-grounded nodes over many speculative ones.

Reference documents:
{docs_text}
"""


def build_graph(board: str) -> dict:
    docs = _all_docs_text(board)
    if not docs:
        raise RuntimeError(f"No documents uploaded to {BOARDS[board]} yet -- upload at least one first.")

    api_key = get_gemini_api_key()
    if not api_key:
        raise RuntimeError("No Gemini API key configured. Add one in Settings.")
    model = get_gemini_model()
    client = genai.Client(api_key=api_key)

    docs_text = "\n\n".join(f"=== {d['filename']} ===\n{d['text']}" for d in docs)[:80_000]
    prompt = WORKSPACE_GRAPH_PROMPT.format(board_label=BOARDS[board], docs_text=docs_text)

    components = usage_module.count_tokens(client, model, {"Reference documents": docs_text})

    t0 = time.time()
    response = client.models.generate_content(
        model=model,
        contents=[prompt],
        config=types.GenerateContentConfig(response_mime_type="application/json", response_schema=WORKSPACE_GRAPH_SCHEMA),
    )
    elapsed = time.time() - t0
    _log_usage(f"workspace-graph:{board}", model, response.usage_metadata, elapsed)
    usage_module.record_call("workspace", board, model, response.usage_metadata, elapsed, components, label="Graph Build")

    graph = json.loads(response.text)
    graph["built_from_doc_ids"] = [d["id"] for d in _load_index(board)]
    _graph_path(board).write_text(json.dumps(graph, indent=2))
    return graph


def get_graph(board: str) -> dict | None:
    p = _graph_path(board)
    if not p.exists():
        return None
    graph = json.loads(p.read_text())
    current_ids = {d["id"] for d in _load_index(board)}
    built_from = set(graph.get("built_from_doc_ids", []))
    graph["stale"] = current_ids != built_from
    return graph


# ---------------- notes ----------------

def _notes_path(board: str) -> Path:
    return _board_dir(board) / "notes.md"


def get_notes(board: str) -> str:
    p = _notes_path(board)
    return p.read_text() if p.exists() else ""


def set_notes(board: str, text: str) -> None:
    _notes_path(board).write_text(text)


def append_notes(board: str, text: str) -> str:
    current = get_notes(board)
    updated = f"{current}\n\n{text}".strip() if current else text
    set_notes(board, updated)
    return updated


def export_docx(board: str) -> Path:
    """Renders the current notes draft (plain text with "- " bullets and blank-line
    paragraph breaks, same lightweight convention the chat UI already uses) into a
    real .docx. Deterministic, no LLM call -- Finalize should be instant."""
    text = get_notes(board)
    doc = Document()
    doc.add_heading(f"{BOARDS[board]} -- Notes", level=0)

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        lines = block.split("\n")
        if all(l.strip().startswith(("- ", "* ")) for l in lines if l.strip()):
            for line in lines:
                line = line.strip()
                if line:
                    doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(block)

    out_path = _board_dir(board) / "notes_export.docx"
    doc.save(str(out_path))
    return out_path
